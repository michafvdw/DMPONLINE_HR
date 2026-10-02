import argparse
import logging
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from dmponline import DMPonline


# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

# Only DMPs using this template will be exported.
TARGET_TEMPLATE = "Hogeschool Rotterdam (HR) Template"

# Default Excel output file.
DEFAULT_OUTPUT_FILE = "DMP_HR_Template_all_questions.xlsx"


# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------


def get_answer_text(question):
    """
    Convert the answer structure returned by DMPonline into readable text.

    Handles:
        - free-text answers
        - scalar answers
        - option-based answers
        - multiple selected options
        - list/tuple answers
        - dictionary answers
    """

    if not question.get("answered"):
        return ""

    answer = question.get("answer")

    if answer is None:
        return ""

    # String / scalar answer
    if isinstance(answer, str):
        return answer

    if isinstance(answer, (int, float, bool)):
        return str(answer)

    # Dictionary answer
    if isinstance(answer, dict):
        # Option-based answers
        options = answer.get("options")

        if options:
            option_texts = []

            for option in options:
                if isinstance(option, dict):
                    text = (
                        option.get("text")
                        or option.get("label")
                        or option.get("name")
                        or option.get("value")
                    )

                    if text is not None:
                        option_texts.append(str(text))
                else:
                    option_texts.append(str(option))

            if option_texts:
                return "\n".join(option_texts)

        # Common text/value fields
        for key in (
            "text",
            "value",
            "answer",
            "content",
            "response",
        ):
            value = answer.get(key)

            if value is not None:
                return str(value)

        # Unknown structure
        return str(answer)

    # List / tuple answer
    if isinstance(answer, (list, tuple)):
        values = []

        for item in answer:
            if isinstance(item, dict):
                value = (
                    item.get("text")
                    or item.get("label")
                    or item.get("name")
                    or item.get("value")
                )

                if value is not None:
                    values.append(str(value))
            else:
                values.append(str(item))

        return "\n".join(values)

    return str(answer)


# ---------------------------------------------------------------------------
# Question extraction
# ---------------------------------------------------------------------------


def extract_all_questions(sections):
    """
    Extract all questions from all sections of a DMP.

    DMPonline can structure questions inside normal sections as well as
    nested sections/subsections. The previous version only looked at:

        section["questions"]

    That can cause the export to stop after an earlier section, for example
    around question 3.1, when later questions are stored in nested sections.

    This function recursively walks through the section hierarchy and collects
    every question found in a ``questions`` list.

    Returns:
        list: Questions in the order in which they occur in the DMP.
    """

    all_questions = []

    # Different DMPonline responses may use different names for nested
    # sections. These are the keys we explicitly recognise as section-like
    # containers.
    nested_section_keys = {
        "sections",
        "subsections",
        "children",
        "question_groups",
        "groups",
    }

    def process_section(section, path=""):
        if not isinstance(section, dict):
            return

        section_number = section.get("number", "")
        section_title = (
            section.get("title")
            or section.get("name")
            or section.get("text")
            or ""
        )

        current_path = path
        if section_number:
            current_path = f"{path}.{section_number}" if path else str(section_number)
        elif section_title:
            current_path = f"{path} > {section_title}" if path else str(section_title)

        # Questions directly contained in this section.
        questions = section.get("questions") or []

        if isinstance(questions, list):
            for question in questions:
                if isinstance(question, dict):
                    # Store a little context on the question so that logging
                    # can tell us exactly where it came from. This does not
                    # change the original question text used as the Excel
                    # header.
                    question_copy = dict(question)
                    question_copy["_section_path"] = current_path
                    all_questions.append(question_copy)

        # Recursively process known nested-section containers.
        for key in nested_section_keys:
            nested = section.get(key)

            if not nested:
                continue

            if isinstance(nested, dict):
                process_section(nested, current_path)

            elif isinstance(nested, list):
                for nested_section in nested:
                    if isinstance(nested_section, dict):
                        process_section(nested_section, current_path)

    if isinstance(sections, list):
        for section in sections:
            process_section(section)
    elif isinstance(sections, dict):
        process_section(sections)

    return all_questions


# ---------------------------------------------------------------------------
# Template helpers
# ---------------------------------------------------------------------------


def get_template_name(plan):
    """
    Try to retrieve the template name from a DMPonline plan.

    DMPonline responses can contain the template information in slightly
    different structures, so this function checks several possibilities.

    Returns:
        str: Template name if found, otherwise "".
    """

    template = plan.get("template")

    if template is None:
        return ""

    # Template is simply a string
    if isinstance(template, str):
        return template.strip()

    # Template is a dictionary
    if isinstance(template, dict):
        for key in (
            "title",
            "name",
            "text",
            "label",
        ):
            value = template.get(key)

            if value:
                return str(value).strip()

        # Sometimes the name may be nested.
        nested_template = template.get("template")

        if isinstance(nested_template, dict):
            for key in (
                "title",
                "name",
                "text",
                "label",
            ):
                value = nested_template.get(key)

                if value:
                    return str(value).strip()

        # Sometimes template information can contain an ID and title
        # inside another object.
        for value in template.values():
            if isinstance(value, dict):
                for key in (
                    "title",
                    "name",
                    "text",
                    "label",
                ):
                    nested_value = value.get(key)

                    if nested_value:
                        return str(nested_value).strip()

    return ""



def plan_uses_target_template(plan, target_template):
    """
    Check whether a DMP uses the requested template.

    Matching is case-insensitive and ignores leading/trailing whitespace.
    """

    template_name = get_template_name(plan)

    if not template_name:
        return False

    return template_name.strip().lower() == target_template.strip().lower()


# ---------------------------------------------------------------------------
# Retrieve all DMPs
# ---------------------------------------------------------------------------


def get_all_plans(api, include_tests=True):
    """
    Retrieve all DMP plans available to the authenticated DMPonline
    organisation administrator.

    The DMPonline v0 API paginates the plans endpoint.
    """

    logging.info("Retrieving all DMPonline plans...")

    all_plans = []
    page = 1

    while True:
        logging.info("Retrieving plans page %s...", page)

        params = {
            "page": page,
            "remove_tests": "false" if include_tests else "true",
        }

        parsed = api.get(
            request="v0/plans",
            params=params,
        )

        if not parsed:
            break

        # Normally DMPonline v0 returns a list.
        if isinstance(parsed, list):
            page_plans = parsed
        elif isinstance(parsed, dict):
            page_plans = (
                parsed.get("data")
                or parsed.get("items")
                or parsed.get("plans")
                or []
            )
        else:
            raise ValueError(
                f"Unexpected response type while retrieving page {page}: "
                f"{type(parsed).__name__}"
            )

        if not page_plans:
            break

        all_plans.extend(page_plans)

        logging.info(
            "Page %s: found %s plans (total so far: %s).",
            page,
            len(page_plans),
            len(all_plans),
        )

        page += 1

    logging.info("Total plans found: %s", len(all_plans))

    return all_plans


# ---------------------------------------------------------------------------
# Filter DMPs by template
# ---------------------------------------------------------------------------


def filter_plans_by_template(plans, target_template):
    """
    Return only plans that use the requested DMPonline template.
    """

    logging.info(
        "Filtering plans using template: %s",
        target_template,
    )

    matching_plans = []
    template_counts = {}

    for plan in plans:
        template_name = get_template_name(plan)

        if template_name:
            template_counts[template_name] = (
                template_counts.get(template_name, 0) + 1
            )

        if plan_uses_target_template(plan, target_template):
            matching_plans.append(plan)

    logging.info(
        "Found %s plans using template '%s'.",
        len(matching_plans),
        target_template,
    )

    logging.info("Templates found in retrieved plans:")

    if template_counts:
        for template_name, count in sorted(template_counts.items()):
            logging.info("  %s: %s plans", template_name, count)
    else:
        logging.warning(
            "No template information was found in the plan responses."
        )

    return matching_plans


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------


def export_all_plans_to_one_excel(plans, api, output_file):
    """
    Export ALL questions from ALL sections of the filtered DMPonline plans
    into ONE Excel file.

    Excel structure:

        DMP Name + ID | 1.1 Question | 1.2 Question | 2.1 Question | ...
        -----------------------------------------------------------------
        DMP A (123)   | Answer        | Answer        | Answer        | ...
        DMP B (456)   | Answer        | Answer        | Answer        | ...

    Each DMP gets one row.
    Each question gets one column.

    Important:
        The question number is NOT added separately. The question text from
        DMPonline is used exactly as the Excel header. This prevents headers
        such as "2 1.2 ..." and keeps them as "1.2 ...".
    """

    logging.info(
        "Creating combined Excel workbook for %s plans...",
        len(plans),
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "All DMPs"

    # -----------------------------------------------------------------------
    # Collect all DMP information
    # -----------------------------------------------------------------------

    all_questions = []
    question_keys = set()
    dmp_data = []

    successful = 0
    failed = 0

    # -----------------------------------------------------------------------
    # Process every DMP
    # -----------------------------------------------------------------------

    for index, plan in enumerate(plans, start=1):
        plan_id = plan.get("id")

        plan_title = plan.get("title") or f"DMP {plan_id}"
        template_name = get_template_name(plan)

        logging.info(
            "Processing plan %s/%s: ID=%s | %s | Template=%s",
            index,
            len(plans),
            plan_id,
            plan_title,
            template_name,
        )

        if not plan_id:
            logging.error("Skipping plan without an ID.")
            failed += 1
            continue

        try:
            # ---------------------------------------------------------------
            # Retrieve the complete DMP
            # ---------------------------------------------------------------

            parsed = api.get(
                request="v0/plans",
                params={
                    "plan": plan_id,
                    "remove_tests": "false",
                },
            )

            if not parsed:
                raise ValueError(
                    f"No DMP content found for plan ID {plan_id}."
                )

            # ---------------------------------------------------------------
            # Extract plan content
            # ---------------------------------------------------------------

            try:
                full_plan = parsed[0]
                plan_content = full_plan["plan_content"][0]
                sections = plan_content["sections"]
            except (IndexError, KeyError, TypeError) as exc:
                raise ValueError(
                    f"Could not find DMP sections for plan {plan_id}."
                ) from exc

            # ---------------------------------------------------------------
            # Extract ALL questions from ALL sections, including nested ones
            # ---------------------------------------------------------------

            questions = extract_all_questions(sections)

            if not questions:
                logging.warning(
                    "Plan %s contains no questions in its section structure.",
                    plan_id,
                )

            actual_plan_title = full_plan.get("title") or plan_title

            # ---------------------------------------------------------------
            # Log the questions found for this DMP
            # ---------------------------------------------------------------

            logging.info(
                "Plan %s: found %s questions across all sections.",
                plan_id,
                len(questions),
            )

            for question in questions:
                question_text = question.get("text", "")

                if question_text is None:
                    question_text = ""

                question_text = str(question_text).strip()

                section_path = question.get("_section_path", "")

                logging.debug(
                    "Plan %s question: [%s] %s",
                    plan_id,
                    section_path,
                    question_text,
                )

            # ---------------------------------------------------------------
            # Store answers for this DMP
            # ---------------------------------------------------------------

            answers_by_question = {}
            dmp_question_count = 0

            for question in questions:
                # -----------------------------------------------------------
                # Question title/text
                # -----------------------------------------------------------

                question_text = question.get("text", "")

                if question_text is None:
                    question_text = ""

                question_text = str(question_text).strip()

                # Ignore question objects without usable text.
                if not question_text:
                    logging.debug(
                        "Plan %s: skipping question without text.",
                        plan_id,
                    )
                    continue

                # -----------------------------------------------------------
                # IMPORTANT:
                # Use the question text itself as the Excel header.
                # Do NOT prepend question["number"].
                #
                # This changes:
                #     2 1.2 Projecttype...
                # into:
                #     1.2 Projecttype...
                # -----------------------------------------------------------

                question_key = question_text

                # -----------------------------------------------------------
                # Answer
                # -----------------------------------------------------------

                answer_text = get_answer_text(question)

                answers_by_question[question_key] = answer_text
                dmp_question_count += 1

                # -----------------------------------------------------------
                # Add question to the global question list
                # -----------------------------------------------------------

                if question_key not in question_keys:
                    question_keys.add(question_key)
                    all_questions.append(question_key)

            # ---------------------------------------------------------------
            # Store the DMP
            # ---------------------------------------------------------------

            dmp_data.append(
                {
                    "plan_id": plan_id,
                    "plan_title": actual_plan_title,
                    "answers": answers_by_question,
                }
            )

            successful += 1

            logging.info(
                "Plan %s processed successfully: %s usable questions.",
                plan_id,
                dmp_question_count,
            )

        except Exception as exc:
            failed += 1

            logging.exception(
                "Could not process plan %s (%s): %s",
                plan_id,
                plan_title,
                exc,
            )

            # Continue processing other DMPs.
            continue

    # -----------------------------------------------------------------------
    # Create the header row
    # -----------------------------------------------------------------------

    worksheet.cell(
        row=1,
        column=1,
        value="DMP Name + ID",
    )

    for column_index, question in enumerate(all_questions, start=2):
        worksheet.cell(
            row=1,
            column=column_index,
            value=question,
        )

    # -----------------------------------------------------------------------
    # Add DMP rows
    # -----------------------------------------------------------------------

    for row_index, dmp in enumerate(dmp_data, start=2):
        plan_id = dmp["plan_id"]
        plan_title = dmp["plan_title"]
        answers = dmp["answers"]

        # Column A = DMP name + ID
        dmp_name_id = f"{plan_title} ({plan_id})"

        worksheet.cell(
            row=row_index,
            column=1,
            value=dmp_name_id,
        )

        # Columns B onward = answers
        for column_index, question in enumerate(all_questions, start=2):
            answer = answers.get(question, "")

            worksheet.cell(
                row=row_index,
                column=column_index,
                value=answer,
            )

    # -----------------------------------------------------------------------
    # Format worksheet
    # -----------------------------------------------------------------------

    format_combined_worksheet(
        worksheet,
        len(all_questions),
    )

    # -----------------------------------------------------------------------
    # Create output directory
    # -----------------------------------------------------------------------

    output_directory = os.path.dirname(os.path.abspath(output_file))
    os.makedirs(output_directory, exist_ok=True)

    # -----------------------------------------------------------------------
    # Save workbook
    # -----------------------------------------------------------------------

    workbook.save(output_file)

    logging.info(
        "Combined Excel report written to: %s",
        os.path.abspath(output_file),
    )

    logging.info(
        "Successfully exported: %s/%s plans.",
        successful,
        len(plans),
    )

    logging.info("Failed: %s plans.", failed)
    logging.info("Total unique question columns: %s", len(all_questions))


# ---------------------------------------------------------------------------
# Excel formatting
# ---------------------------------------------------------------------------


def format_combined_worksheet(ws, number_of_questions):
    """
    Format the horizontally structured worksheet.
    """

    # -----------------------------------------------------------------------
    # Styles
    # -----------------------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    dmp_font = Font(
        bold=True,
    )

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # -----------------------------------------------------------------------
    # Header row
    # -----------------------------------------------------------------------

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    ws.row_dimensions[1].height = 100

    # -----------------------------------------------------------------------
    # Data rows
    # -----------------------------------------------------------------------

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = thin_border

        # DMP name + ID is bold.
        row[0].font = dmp_font

    # -----------------------------------------------------------------------
    # Column widths
    # -----------------------------------------------------------------------

    ws.column_dimensions["A"].width = 40

    for column_index in range(2, number_of_questions + 2):
        column_letter = ws.cell(
            row=1,
            column=column_index,
        ).column_letter

        ws.column_dimensions[column_letter].width = 35

    # -----------------------------------------------------------------------
    # Row heights
    # -----------------------------------------------------------------------

    for row_number in range(2, ws.max_row + 1):
        ws.row_dimensions[row_number].height = 100

    # -----------------------------------------------------------------------
    # Freeze panes
    # -----------------------------------------------------------------------

    ws.freeze_panes = "B2"

    # -----------------------------------------------------------------------
    # Auto filter
    # -----------------------------------------------------------------------

    if ws.max_row >= 2 and number_of_questions >= 1:
        last_column = ws.cell(
            row=1,
            column=number_of_questions + 1,
        ).column_letter

        ws.auto_filter.ref = f"A1:{last_column}{ws.max_row}"


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Export all questions from all sections of DMPonline plans "
            f"using the '{TARGET_TEMPLATE}' template into ONE Excel file."
        )
    )

    # -----------------------------------------------------------------------
    # API token
    # -----------------------------------------------------------------------

    parser.add_argument(
        "-t",
        "--token",
        dest="token",
        required=True,
        help="DMPonline API token",
    )

    # -----------------------------------------------------------------------
    # User email
    # -----------------------------------------------------------------------

    parser.add_argument(
        "-u",
        "--user",
        dest="token_user",
        required=True,
        help="DMPonline user email address",
    )

    # -----------------------------------------------------------------------
    # Output file
    # -----------------------------------------------------------------------

    parser.add_argument(
        "-o",
        "--output",
        dest="output_file",
        default=DEFAULT_OUTPUT_FILE,
        help=(
            "Excel file to create. "
            f"Default: {DEFAULT_OUTPUT_FILE}"
        ),
    )

    # -----------------------------------------------------------------------
    # Exclude test plans
    # -----------------------------------------------------------------------

    parser.add_argument(
        "--exclude-tests",
        action="store_true",
        help="Exclude DMPonline test plans.",
    )

    args = parser.parse_args()

    # -----------------------------------------------------------------------
    # Logging
    # -----------------------------------------------------------------------

    logging.basicConfig(
        level=logging.INFO,
        format="%(levelname)s: %(message)s",
    )

    # -----------------------------------------------------------------------
    # Validate output extension
    # -----------------------------------------------------------------------

    if os.path.splitext(args.output_file)[1].lower() != ".xlsx":
        raise ValueError("The output file must have an .xlsx extension.")

    # -----------------------------------------------------------------------
    # Create DMPonline API client
    # -----------------------------------------------------------------------

    logging.info("Connecting to DMPonline...")

    api = DMPonline(
        token=args.token,
        token_user=args.token_user,
    )

    # -----------------------------------------------------------------------
    # Get all plans
    # -----------------------------------------------------------------------

    plans = get_all_plans(
        api=api,
        include_tests=not args.exclude_tests,
    )

    if not plans:
        logging.warning("No plans were returned by the DMPonline API.")
        return

    # -----------------------------------------------------------------------
    # Filter by template
    # -----------------------------------------------------------------------

    filtered_plans = filter_plans_by_template(
        plans=plans,
        target_template=TARGET_TEMPLATE,
    )

    if not filtered_plans:
        logging.warning(
            "No DMPs were found using the template: %s",
            TARGET_TEMPLATE,
        )
        return

    # -----------------------------------------------------------------------
    # Export filtered plans into ONE workbook
    # -----------------------------------------------------------------------

    export_all_plans_to_one_excel(
        plans=filtered_plans,
        api=api,
        output_file=args.output_file,
    )


# ---------------------------------------------------------------------------
# Program entry point
# ---------------------------------------------------------------------------


if __name__ == "__main__":
    main()
