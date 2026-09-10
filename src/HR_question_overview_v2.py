import argparse
import logging
import os

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from dmponline import DMPonline


# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------

def get_answer_text(question):
    """
    Convert the answer structure returned by DMPonline into readable text.

    Handles:
        - free-text answers
        - scalar answers
        - option-based answers
        - multiple selected options
        - list/tuple answers
        - dictionary answers
    """

    # If DMPonline says the question was not answered,
    # return an empty string rather than removing the question.
    if not question.get("answered"):
        return ""

    answer = question.get("answer")

    if answer is None:
        return ""

    # -----------------------------------------------------------------------
    # String / scalar answer
    # -----------------------------------------------------------------------

    if isinstance(answer, str):
        return answer

    if isinstance(answer, (int, float, bool)):
        return str(answer)

    # -----------------------------------------------------------------------
    # Dictionary answer
    # -----------------------------------------------------------------------

    if isinstance(answer, dict):

        # Option-based answers
        options = answer.get("options")

        if options:
            option_texts = []

            for option in options:

                if isinstance(option, dict):
                    text = (
                        option.get("text")
                        or option.get("label")
                        or option.get("name")
                        or option.get("value")
                    )

                    if text is not None:
                        option_texts.append(str(text))

                else:
                    option_texts.append(str(option))

            if option_texts:
                return "\n".join(option_texts)

        # Common text/value fields
        for key in (
            "text",
            "value",
            "answer",
            "content",
            "response"
        ):
            value = answer.get(key)

            if value is not None:
                return str(value)

        # If the structure is unknown, return it as text
        # so that information is not silently lost.
        return str(answer)

    # -----------------------------------------------------------------------
    # List / tuple answer
    # -----------------------------------------------------------------------

    if isinstance(answer, (list, tuple)):

        values = []

        for item in answer:

            if isinstance(item, dict):

                value = (
                    item.get("text")
                    or item.get("label")
                    or item.get("name")
                    or item.get("value")
                )

                if value is not None:
                    values.append(str(value))
                else:
                    values.append(str(item))

            else:
                values.append(str(item))

        return "\n".join(values)

    # -----------------------------------------------------------------------
    # Fallback
    # -----------------------------------------------------------------------

    return str(answer)


def find_section(sections, section_number):
    """
    Find a section by its section number.
    """

    for section in sections:

        number = section.get("number")

        # Handle both integer and string section numbers.
        if str(number) == str(section_number):
            return section

    return None


def format_worksheet(ws):
    """
    Apply formatting to the Excel worksheet.
    """

    # -----------------------------------------------------------------------
    # Styles
    # -----------------------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    question_number_font = Font(
        bold=True
    )

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # -----------------------------------------------------------------------
    # Header row
    # -----------------------------------------------------------------------

    for cell in ws[1]:

        cell.fill = header_fill
        cell.font = header_font

        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border

    ws.row_dimensions[1].height = 30

    # -----------------------------------------------------------------------
    # Data rows
    # -----------------------------------------------------------------------

    for row in ws.iter_rows(min_row=2):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            cell.border = thin_border

    # -----------------------------------------------------------------------
    # Make question number bold
    # -----------------------------------------------------------------------

    for row in range(2, ws.max_row + 1):

        ws.cell(
            row=row,
            column=1
        ).font = question_number_font

    # -----------------------------------------------------------------------
    # Column widths
    # -----------------------------------------------------------------------

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 80
    ws.column_dimensions["C"].width = 100

    # -----------------------------------------------------------------------
    # Row heights
    # -----------------------------------------------------------------------

    for row in range(2, ws.max_row + 1):

        # Give every answer enough room to be readable.
        ws.row_dimensions[row].height = 75

    # -----------------------------------------------------------------------
    # Freeze header row
    # -----------------------------------------------------------------------

    ws.freeze_panes = "A2"

    # -----------------------------------------------------------------------
    # Filter
    # -----------------------------------------------------------------------

    if ws.max_row >= 2:

        ws.auto_filter.ref = (
            f"A1:C{ws.max_row}"
        )


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def question_overview(
    plan_id,
    api,
    section_number,
    output_file
):
    """
    Export all questions and answers from one DMPonline section.

    The Excel file contains only:

        Question number
        Question
        Answer

    Unanswered questions are included with an empty answer cell.
    """

    logging.info(
        f"Retrieving DMP {plan_id}..."
    )

    # -----------------------------------------------------------------------
    # Get raw DMP response
    # -----------------------------------------------------------------------

    parsed = api.get(
        request=f"v0/plans?plan={plan_id}",
        params={
            "remove_tests": "false"
        }
    )

    if not parsed:

        raise ValueError(
            f"No DMP found for plan ID {plan_id}."
        )

    # -----------------------------------------------------------------------
    # Extract DMP content
    # -----------------------------------------------------------------------

    try:

        plan = parsed[0]

        plan_content = plan["plan_content"][0]

        sections = plan_content["sections"]

    except (IndexError, KeyError, TypeError) as exc:

        raise ValueError(
            "Could not find the DMP sections in the API response."
        ) from exc

    # -----------------------------------------------------------------------
    # Find requested section
    # -----------------------------------------------------------------------

    selected_section = find_section(
        sections,
        section_number
    )

    if selected_section is None:

        available_sections = [
            str(section.get("number"))
            for section in sections
        ]

        raise ValueError(
            f"Section {section_number} was not found. "
            f"Available sections: {', '.join(available_sections)}"
        )

    section_title = (
        selected_section.get("title")
        or selected_section.get("name")
        or f"Section {section_number}"
    )

    questions = (
        selected_section.get("questions")
        or []
    )

    logging.info(
        f"Found section {section_number}: {section_title}"
    )

    logging.info(
        f"Found {len(questions)} questions."
    )

    # -----------------------------------------------------------------------
    # Create Excel workbook
    # -----------------------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        f"Section {section_number}"
    )

    # -----------------------------------------------------------------------
    # Headers
    # -----------------------------------------------------------------------

    worksheet["A1"] = "Question number"
    worksheet["B1"] = "Question"
    worksheet["C1"] = "Answer"

    # -----------------------------------------------------------------------
    # Add questions and answers
    # -----------------------------------------------------------------------

    row = 2

    for question in questions:

        question_number = question.get(
            "number",
            ""
        )

        question_text = question.get(
            "text",
            ""
        )

        answer_text = get_answer_text(
            question
        )

        worksheet.cell(
            row=row,
            column=1,
            value=question_number
        )

        worksheet.cell(
            row=row,
            column=2,
            value=question_text
        )

        worksheet.cell(
            row=row,
            column=3,
            value=answer_text
        )

        row += 1

    # -----------------------------------------------------------------------
    # Format worksheet
    # -----------------------------------------------------------------------

    format_worksheet(
        worksheet
    )

    # -----------------------------------------------------------------------
    # Validate output filename
    # -----------------------------------------------------------------------

    if not output_file:

        raise ValueError(
            "An output Excel filename is required."
        )

    extension = os.path.splitext(
        output_file
    )[-1].lower()

    if extension != ".xlsx":

        raise ValueError(
            "The output file must have an .xlsx extension."
        )

    # -----------------------------------------------------------------------
    # Save workbook
    # -----------------------------------------------------------------------

    workbook.save(
        output_file
    )

    logging.info(
        f"Excel report written to: {output_file}"
    )


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():

    parser = argparse.ArgumentParser(
        description=(
            "Export all questions and answers from one "
            "DMPonline section to Excel."
        )
    )

    # -----------------------------------------------------------------------
    # DMP ID
    # -----------------------------------------------------------------------

    parser.add_argument(
        "-i",
        "--id",
        dest="plan_id",
        required=True,
        help="DMPonline plan ID"
    )

    # -----------------------------------------------------------------------
    # API token
    # -----------------------------------------------------------------------

    parser.add_argument(
        "-t",
        "--token",
        dest="token",
        required=True,
        help="DMPonline API token"
    )

    # -----------------------------------------------------------------------
    # User email
    # -----------------------------------------------------------------------

    parser.add_argument(
        "-u",
        "--user",
        dest="token_user",
        required=True,
        help="DMPonline user email address"
    )

    # -----------------------------------------------------------------------
    # Section number
    # -----------------------------------------------------------------------

    parser.add_argument(
        "-s",
        "--section",
        dest="section_number",
        required=True,
        help="Section number to export, e.g. 1, 2, 3, 4 or 5"
    )

    # -----------------------------------------------------------------------
    # Output file
    # -----------------------------------------------------------------------

    parser.add_argument(
        "-o",
        "--output",
        dest="output_file",
        required=True,
        help="Output Excel filename, e.g. DMP_section_2.xlsx"
    )

    # -----------------------------------------------------------------------
    # Parse arguments
    # -----------------------------------------------------------------------

    args = parser.parse_args()

    # -----------------------------------------------------------------------
    # Logging
    # -----------------------------------------------------------------------

    logging.basicConfig(
        level=logging.INFO,
        format="%(levelname)s: %(message)s"
    )

    # -----------------------------------------------------------------------
    # Create DMPonline API client
    # -----------------------------------------------------------------------

    logging.info(
        "Connecting to DMPonline..."
    )

    api = DMPonline(
        token=args.token,
        token_user=args.token_user
    )

    # -----------------------------------------------------------------------
    # Export section
    # -----------------------------------------------------------------------

    question_overview(
        plan_id=args.plan_id,
        api=api,
        section_number=args.section_number,
        output_file=args.output_file
    )


# ---------------------------------------------------------------------------
# Program entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    main()

