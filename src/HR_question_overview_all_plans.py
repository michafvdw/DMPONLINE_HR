import argparse
import logging
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from dmponline import DMPonline


# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------

def get_answer_text(question):
    """
    Convert the answer structure returned by DMPonline into readable text.

    Handles:
        - free-text answers
        - scalar answers
        - option-based answers
        - multiple selected options
        - list/tuple answers
        - dictionary answers
    """

    if not question.get("answered"):
        return ""

    answer = question.get("answer")

    if answer is None:
        return ""

    # String / scalar answer
    if isinstance(answer, str):
        return answer

    if isinstance(answer, (int, float, bool)):
        return str(answer)

    # Dictionary answer
    if isinstance(answer, dict):

        # Option-based answers
        options = answer.get("options")

        if options:
            option_texts = []

            for option in options:
                if isinstance(option, dict):
                    text = (
                        option.get("text")
                        or option.get("label")
                        or option.get("name")
                        or option.get("value")
                    )

                    if text is not None:
                        option_texts.append(str(text))
                else:
                    option_texts.append(str(option))

            if option_texts:
                return "\n".join(option_texts)

        # Common text/value fields
        for key in (
            "text",
            "value",
            "answer",
            "content",
            "response",
        ):
            value = answer.get(key)

            if value is not None:
                return str(value)

        # Unknown structure: keep the information instead of losing it.
        return str(answer)

    # List / tuple answer
    if isinstance(answer, (list, tuple)):

        values = []

        for item in answer:
            if isinstance(item, dict):
                value = (
                    item.get("text")
                    or item.get("label")
                    or item.get("name")
                    or item.get("value")
                )

                if value is not None:
                    values.append(str(value))
                else:
                    values.append(str(item))
            else:
                values.append(str(item))

        return "\n".join(values)

    return str(answer)


def find_section(sections, section_number):
    """Find a section by its section number."""

    for section in sections:
        number = section.get("number")

        if str(number) == str(section_number):
            return section

    return None


def get_all_plans(api, include_tests=True):
    """
    Retrieve all DMP plans available to the authenticated DMPonline
    organisation administrator.

    The DMPonline v0 API paginates the plans endpoint. The current API
    returns 20 plans per page, so we continue requesting pages until an
    empty page is returned.

    Returns:
        list: A list of plan metadata/content dictionaries.
    """

    logging.info("Retrieving all DMPonline plans...")

    all_plans = []
    page = 1

    while True:
        logging.info("Retrieving plans page %s...", page)

        params = {
            "page": page,
            "remove_tests": "false" if include_tests else "true",
        }

        parsed = api.get(
            request="v0/plans",
            params=params,
        )

        if not parsed:
            break

        # Normally DMPonline v0 returns a list.
        # This also handles a dictionary containing a common data/items key,
        # in case the API wrapper returns a paginated structure.
        if isinstance(parsed, list):
            page_plans = parsed

        elif isinstance(parsed, dict):
            page_plans = (
                parsed.get("data")
                or parsed.get("items")
                or parsed.get("plans")
                or []
            )

        else:
            raise ValueError(
                f"Unexpected response type while retrieving page {page}: "
                f"{type(parsed).__name__}"
            )

        if not page_plans:
            break

        all_plans.extend(page_plans)

        logging.info(
            "Page %s: found %s plans (total so far: %s).",
            page,
            len(page_plans),
            len(all_plans),
        )

        page += 1

    logging.info("Total plans found: %s", len(all_plans))

    return all_plans


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------

def export_all_plans_to_one_excel(
    plans,
    api,
    section_number,
    output_file,
):
    """
    Export the selected section from ALL DMPonline plans into ONE Excel file.

    Every row represents one question from one DMP.

    Columns:
        Plan ID
        Plan title
        Question number
        Question
        Answer
    """

    logging.info(
        "Creating combined Excel workbook for %s plans...",
        len(plans),
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "All DMPs"

    # Headers
    worksheet["A1"] = "Plan ID"
    worksheet["B1"] = "Plan title"
    worksheet["C1"] = "Question number"
    worksheet["D1"] = "Question"
    worksheet["E1"] = "Answer"

    row = 2
    successful = 0
    failed = 0

    for index, plan in enumerate(plans, start=1):

        plan_id = plan.get("id")
        plan_title = plan.get("title") or f"DMP {plan_id}"

        logging.info(
            "Processing plan %s/%s: ID=%s | %s",
            index,
            len(plans),
            plan_id,
            plan_title,
        )

        if not plan_id:
            logging.error("Skipping plan without an ID.")
            failed += 1
            continue

        try:
            # Retrieve the complete DMP.
            parsed = api.get(
                request="v0/plans",
                params={
                    "plan": plan_id,
                    "remove_tests": "false",
                },
            )

            if not parsed:
                raise ValueError(
                    f"No DMP content found for plan ID {plan_id}."
                )

            try:
                full_plan = parsed[0]
                plan_content = full_plan["plan_content"][0]
                sections = plan_content["sections"]

            except (IndexError, KeyError, TypeError) as exc:
                raise ValueError(
                    f"Could not find DMP sections for plan {plan_id}."
                ) from exc

            selected_section = find_section(
                sections,
                section_number,
            )

            if selected_section is None:
                available_sections = [
                    str(section.get("number"))
                    for section in sections
                ]

                raise ValueError(
                    f"Section {section_number} was not found. "
                    f"Available sections: "
                    f"{', '.join(available_sections)}"
                )

            questions = selected_section.get("questions") or []

            actual_plan_title = (
                full_plan.get("title")
                or plan_title
            )

            # Add every question from this DMP to the same worksheet.
            for question in questions:

                question_number = question.get(
                    "number",
                    "",
                )

                question_text = question.get(
                    "text",
                    "",
                )

                answer_text = get_answer_text(question)

                worksheet.cell(
                    row=row,
                    column=1,
                    value=plan_id,
                )

                worksheet.cell(
                    row=row,
                    column=2,
                    value=actual_plan_title,
                )

                worksheet.cell(
                    row=row,
                    column=3,
                    value=question_number,
                )

                worksheet.cell(
                    row=row,
                    column=4,
                    value=question_text,
                )

                worksheet.cell(
                    row=row,
                    column=5,
                    value=answer_text,
                )

                row += 1

            successful += 1

            logging.info(
                "Plan %s exported: %s questions.",
                plan_id,
                len(questions),
            )

        except Exception as exc:
            failed += 1

            logging.error(
                "Could not export plan %s (%s): %s",
                plan_id,
                plan_title,
                exc,
            )

            # Continue with the next DMP.
            continue

    # -----------------------------------------------------------------------
    # Formatting
    # -----------------------------------------------------------------------

    format_combined_worksheet(worksheet)

    # Make sure the output directory exists.
    os.makedirs(
        os.path.dirname(os.path.abspath(output_file)),
        exist_ok=True,
    )

    workbook.save(output_file)

    logging.info(
        "Combined Excel report written to: %s",
        os.path.abspath(output_file),
    )

    logging.info(
        "Successfully exported: %s/%s plans.",
        successful,
        len(plans),
    )

    logging.info(
        "Failed: %s plans.",
        failed,
    )


def format_combined_worksheet(ws):
    """
    Format the combined worksheet.
    """

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    question_number_font = Font(
        bold=True,
    )

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # Header row
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True,
        )
        cell.border = thin_border

    ws.row_dimensions[1].height = 30

    # Data rows
    for row in ws.iter_rows(min_row=2):

        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

            cell.border = thin_border

        # Make question number bold.
        row[2].font = question_number_font

    # Column widths
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 80
    ws.column_dimensions["E"].width = 100

    # Row heights
    for row_number in range(2, ws.max_row + 1):
        ws.row_dimensions[row_number].height = 75

    # Freeze header
    ws.freeze_panes = "A2"

    # Filter
    if ws.max_row >= 2:
        ws.auto_filter.ref = f"A1:E{ws.max_row}"


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def main():

    parser = argparse.ArgumentParser(
        description=(
            "Export a selected section from ALL DMPonline plans "
            "into ONE Excel file."
        )
    )

    # API token
    parser.add_argument(
        "-t",
        "--token",
        dest="token",
        required=True,
        help="DMPonline API token",
    )

    # User email
    parser.add_argument(
        "-u",
        "--user",
        dest="token_user",
        required=True,
        help="DMPonline user email address",
    )

    # Section number
    parser.add_argument(
        "-s",
        "--section",
        dest="section_number",
        required=True,
        help="Section number to export, e.g. 1, 2, 3, 4 or 5",
    )

    # Output file
    parser.add_argument(
        "-o",
        "--output",
        dest="output_file",
        default="DMP_all_plans.xlsx",
        help=(
            "Excel file to create. "
            "Default: DMP_all_plans.xlsx"
        ),
    )

    # Whether to exclude test plans.
    parser.add_argument(
        "--exclude-tests",
        action="store_true",
        help="Exclude DMPonline test plans.",
    )

    args = parser.parse_args()

    # Logging
    logging.basicConfig(
        level=logging.INFO,
        format="%(levelname)s: %(message)s",
    )

    # Validate output extension.
    if os.path.splitext(args.output_file)[1].lower() != ".xlsx":
        raise ValueError(
            "The output file must have an .xlsx extension."
        )

    # Create DMPonline API client.
    logging.info("Connecting to DMPonline...")

    api = DMPonline(
        token=args.token,
        token_user=args.token_user,
    )

    # Get all plans.
    plans = get_all_plans(
        api=api,
        include_tests=not args.exclude_tests,
    )

    if not plans:
        logging.warning(
            "No plans were returned by the DMPonline API."
        )
        return

    # Export all plans into ONE workbook.
    export_all_plans_to_one_excel(
        plans=plans,
        api=api,
        section_number=args.section_number,
        output_file=args.output_file,
    )


# ---------------------------------------------------------------------------
# Program entry point
# ---------------------------------------------------------------------------

if __name__ == "__main__":
    main()
