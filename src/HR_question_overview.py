
import argparse
import logging
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from dmponline import DMPonline


# ---------------------------------------------------------------------------
# Helper functions
# ---------------------------------------------------------------------------

def get_answer_text(question):
    """
    Convert the answer structure returned by DMPonline into readable text.

    DMPonline can return different answer structures depending on the
    question type. This function tries to handle the common cases:
      - free-text answers
      - option-based answers
      - multiple selected options
      - simple scalar answers
    """

    if not question.get('answered'):
        return None

    answer = question.get('answer')

    if answer is None:
        return ""

    # ---------------------------------------------------------------
    # String / scalar answer
    # ---------------------------------------------------------------

    if isinstance(answer, str):
        return answer

    if isinstance(answer, (int, float, bool)):
        return str(answer)

    # ---------------------------------------------------------------
    # Dictionary answer
    # ---------------------------------------------------------------

    if isinstance(answer, dict):

        # Option-based answers
        #
        # Example:
        # {
        #     "options": [
        #         {"text": "Yes"},
        #         {"text": "No"}
        #     ]
        # }
        options = answer.get('options')

        if options:
            option_texts = []

            for option in options:

                if isinstance(option, dict):
                    text = (
                        option.get('text')
                        or option.get('label')
                        or option.get('name')
                    )

                    if text:
                        option_texts.append(str(text))

                else:
                    option_texts.append(str(option))

            if option_texts:
                return "\n".join(option_texts)

        # Common text/value fields
        for key in (
            'text',
            'value',
            'answer',
            'content',
            'response'
        ):
            value = answer.get(key)

            if value is not None:
                return str(value)

        # -----------------------------------------------------------
        # If we don't recognise the structure, return it as text
        # rather than silently losing the answer.
        # -----------------------------------------------------------

        return str(answer)

    # ---------------------------------------------------------------
    # List / tuple answer
    # ---------------------------------------------------------------

    if isinstance(answer, (list, tuple)):

        values = []

        for item in answer:

            if isinstance(item, dict):
                value = (
                    item.get('text')
                    or item.get('label')
                    or item.get('name')
                    or item.get('value')
                )

                if value is not None:
                    values.append(str(value))
                else:
                    values.append(str(item))

            else:
                values.append(str(item))

        return "\n".join(values)

    # ---------------------------------------------------------------
    # Fallback
    # ---------------------------------------------------------------

    return str(answer)


def safe_sheet_name(name, used_names):
    """
    Create an Excel-compatible worksheet name.

    Excel:
      - limits sheet names to 31 characters
      - does not allow: \\ / * ? : [ ]
      - does not allow duplicate sheet names
    """

    if not name:
        name = "Section"

    # Remove invalid Excel characters
    invalid_characters = ['\\', '/', '*', '?', ':', '[', ']']

    for character in invalid_characters:
        name = name.replace(character, '-')

    name = name.strip()

    if not name:
        name = "Section"

    # Excel limits worksheet names to 31 characters
    name = name[:31]

    original_name = name
    counter = 2

    while name in used_names:

        suffix = f" ({counter})"

        name = original_name[:31 - len(suffix)] + suffix

        counter += 1

    used_names.add(name)

    return name


def format_worksheet(ws):
    """
    Apply formatting to a worksheet.
    """

    # ---------------------------------------------------------------
    # Colours / styles
    # ---------------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78"
    )

    section_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    question_font = Font(
        bold=True
    )

    title_font = Font(
        bold=True,
        size=16
    )

    subtitle_font = Font(
        bold=True,
        size=11
    )

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # ---------------------------------------------------------------
    # Title
    # ---------------------------------------------------------------

    ws["A1"].font = title_font

    # ---------------------------------------------------------------
    # Header row
    # ---------------------------------------------------------------

    header_row = 4

    for cell in ws[header_row]:

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
            wrap_text=True
        )

        cell.border = thin_border

    ws.row_dimensions[header_row].height = 30

    # ---------------------------------------------------------------
    # Data rows
    # ---------------------------------------------------------------

    for row in ws.iter_rows(
        min_row=header_row + 1
    ):

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

            cell.border = thin_border

    # ---------------------------------------------------------------
    # Make question column bold
    # ---------------------------------------------------------------

    for row in range(
        header_row + 1,
        ws.max_row + 1
    ):

        # Question text is column C
        if ws.cell(row=row, column=3).value:
            ws.cell(
                row=row,
                column=3
            ).font = question_font

    # ---------------------------------------------------------------
    # Column widths
    # ---------------------------------------------------------------

    widths = {
        "A": 14,   # Question
        "B": 15,   # Type
        "C": 65,   # Question text
        "D": 100   # Answer
    }

    for column, width in widths.items():
        ws.column_dimensions[column].width = width

    # ---------------------------------------------------------------
    # Freeze question headers
    # ---------------------------------------------------------------

    ws.freeze_panes = "A5"

    # ---------------------------------------------------------------
    # Autofilter
    # ---------------------------------------------------------------

    if ws.max_row >= 4:
        ws.auto_filter.ref = (
            f"A4:D{ws.max_row}"
        )


# ---------------------------------------------------------------------------
# Main Excel export
# ---------------------------------------------------------------------------

def question_overview(plan_id, api, output_file=None):
    """
    Export the answered questions of a DMP to an Excel workbook.

    Each DMP section gets its own worksheet.

    The workbook also contains an Overview sheet.
    """

    # ---------------------------------------------------------------
    # Retrieve the DMP
    # ---------------------------------------------------------------

    plan = api.get_plan_v0(plan_id)

    if plan is None:
        raise ValueError(
            f"No DMP found for plan ID {plan_id}"
        )

    # ---------------------------------------------------------------
    # Get plan content
    #
    # get_plan_v0() in the current project returns the plan content
    # structure used by question_overview.py.
    # ---------------------------------------------------------------

    try:
        plan_content = plan.plan_content
    except AttributeError:

        # Some versions / API responses may return a dictionary
        # rather than an object.
        if isinstance(plan, dict):
            plan_content = plan.get("plan_content")
        else:
            raise ValueError(
                "Could not find 'plan_content' in the DMP response."
            )

    if not plan_content:
        raise ValueError(
            f"DMP {plan_id} does not contain any plan content."
        )

    # ---------------------------------------------------------------
    # Create workbook
    # ---------------------------------------------------------------

    workbook = Workbook()

    # Remove default sheet
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_sheet_names = set()

    # ---------------------------------------------------------------
    # Overview worksheet
    # ---------------------------------------------------------------

    overview = workbook.create_sheet("Overview")
    used_sheet_names.add("Overview")

    overview["A1"] = "DMPonline DMP overview"
    overview["A1"].font = Font(
        bold=True,
        size=18
    )

    overview["A3"] = "Plan ID"
    overview["B3"] = plan_id

    overview["A5"] = "Sections"
    overview["B5"] = 0

    overview["A7"] = "Section"
    overview["B7"] = "Worksheet"
    overview["C7"] = "Answered questions"

    for cell in overview[7]:

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78"
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.alignment = Alignment(
            vertical="center",
            wrap_text=True
        )

    overview_row = 8

    section_count = 0

    # ---------------------------------------------------------------
    # Extract sections
    # ---------------------------------------------------------------

    # The existing project accesses:
    #
    # plan.plan_content[0][0]['sections']
    #
    # Keep that structure for compatibility with the current API
    # implementation.

    try:
        sections = plan_content[0][0]["sections"]

    except (KeyError, IndexError, TypeError):

        # Try a slightly less nested structure as a fallback.
        try:
            sections = plan_content[0]["sections"]
        except (KeyError, IndexError, TypeError):
            raise ValueError(
                "Could not find sections in the DMP plan content."
            )

    for section in sections:

        section_number = section.get(
            "number",
            section_count + 1
        )

        section_title = (
            section.get("title")
            or section.get("name")
            or f"Section {section_number}"
        )

        questions = section.get(
            "questions",
            []
        )

        # -----------------------------------------------------------
        # Keep only answered questions
        # -----------------------------------------------------------

        answered_questions = []

        for question in questions:

            if not question.get("answered"):
                continue

            answer_text = get_answer_text(question)

            # Do not throw away answers that are explicitly empty.
            if answer_text is None:
                continue

            answered_questions.append(
                (
                    question,
                    answer_text
                )
            )

        # -----------------------------------------------------------
        # Create worksheet
        # -----------------------------------------------------------

        sheet_name = safe_sheet_name(
            f"{section_number} - {section_title}",
            used_sheet_names
        )

        ws = workbook.create_sheet(sheet_name)

        # -----------------------------------------------------------
        # Section title
        # -----------------------------------------------------------

        ws["A1"] = (
            f"Section {section_number}: "
            f"{section_title}"
        )

        ws["A1"].font = Font(
            bold=True,
            size=16
        )

        ws.merge_cells(
            start_row=1,
            start_column=1,
            end_row=1,
            end_column=4
        )

        # -----------------------------------------------------------
        # Number of answered questions
        # -----------------------------------------------------------

        ws["A2"] = "Answered questions"
        ws["B2"] = len(answered_questions)

        # -----------------------------------------------------------
        # Table headers
        # -----------------------------------------------------------

        ws["A4"] = "Question"
        ws["B4"] = "Type"
        ws["C4"] = "Question text"
        ws["D4"] = "Answer"

        # -----------------------------------------------------------
        # Questions
        # -----------------------------------------------------------

        row = 5

        for question, answer_text in answered_questions:

            question_number = question.get(
                "number",
                ""
            )

            question_type = question.get(
                "format",
                ""
            )

            question_text = question.get(
                "text",
                ""
            )

            ws.cell(
                row=row,
                column=1,
                value=question_number
            )

            ws.cell(
                row=row,
                column=2,
                value=question_type
            )

            ws.cell(
                row=row,
                column=3,
                value=question_text
            )

            ws.cell(
                row=row,
                column=4,
                value=answer_text
            )

            row += 1

        # -----------------------------------------------------------
        # Format worksheet
        # -----------------------------------------------------------

        format_worksheet(ws)

        # -----------------------------------------------------------
        # Add section to overview
        # -----------------------------------------------------------

        overview.cell(
            row=overview_row,
            column=1,
            value=f"{section_number}: {section_title}"
        )

        overview.cell(
            row=overview_row,
            column=2,
            value=sheet_name
        )

        overview.cell(
            row=overview_row,
            column=3,
            value=len(answered_questions)
        )

        overview_row += 1
        section_count += 1

    # ---------------------------------------------------------------
    # Finish Overview sheet
    # ---------------------------------------------------------------

    overview["B5"] = section_count

    for row in overview.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )

    for cell in overview[7]:

        cell.border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

    overview.column_dimensions["A"].width = 55
    overview.column_dimensions["B"].width = 35
    overview.column_dimensions["C"].width = 22

    overview.freeze_panes = "A8"

    if overview_row > 8:
        overview.auto_filter.ref = (
            f"A7:C{overview_row - 1}"
        )

    # ---------------------------------------------------------------
    # Save workbook
    # ---------------------------------------------------------------

    if output_file is not None:

        extension = os.path.splitext(
            output_file
        )[-1].lower()

        if extension != ".xlsx":
            raise ValueError(
                "This report format requires an .xlsx output file."
            )

        workbook.save(output_file)

        logging.info(
            f"Excel report written to {output_file}"
        )

    else:

        logging.warning(
            "No output file specified. "
            "Use -o filename.xlsx"
        )


# ---------------------------------------------------------------------------
# Command line interface
# ---------------------------------------------------------------------------

def main():

    logging.basicConfig(
        stream=sys.stdout,
        level=logging.INFO
    )

    parser = argparse.ArgumentParser(
        description=(
            "Export answered questions from a specific "
            "DMPonline plan to an Excel workbook."
        )
    )

    # Required arguments
    required_named = parser.add_argument_group(
        "required named arguments"
    )

    required_named.add_argument(
        "-i",
        "--dmponline-plan-id",
        type=int,
        help="DMPonline plan ID.",
        required=True
    )

    required_named.add_argument(
        "-t",
        "--dmponline-api-token",
        help="DMPonline API access token.",
        required=True
    )

    # Optional arguments

    parser.add_argument(
        "-u",
        "--dmponline-user-email",
        default=None,
        help=(
            "Username (email) of user corresponding to "
            "DMPONLINE_API_TOKEN."
        )
    )

    parser.add_argument(
        "--do-not-verify",
        action="store_true",
        help=(
            "Disable SSL certificate verification."
        )
    )

    parser.add_argument(
        "-c",
        "--cert-file",
        default=None,
        help="SSL certificate file."
    )

    parser.add_argument(
        "-o",
        "--output-file",
        default=None,
        required=True,
        help=(
            "Output Excel file, e.g. dmp_12345.xlsx"
        )
    )

    args = parser.parse_args()

    # ---------------------------------------------------------------
    # SSL verification
    # ---------------------------------------------------------------

    if (
        args.cert_file is not None
        and os.path.exists(args.cert_file)
    ):
        verify = args.cert_file

    else:
        verify = not args.do_not_verify

    logging.debug(
        f"Calling API with verify={verify}"
    )

    # ---------------------------------------------------------------
    # Initialise API
    # ---------------------------------------------------------------

    dmp_api = DMPonline(
        args.dmponline_api_token,
        verify=verify,
        token_user=args.dmponline_user_email
    )

    # ---------------------------------------------------------------
    # Generate report
    # ---------------------------------------------------------------

    question_overview(
        args.dmponline_plan_id,
        dmp_api,
        output_file=args.output_file
    )


if __name__ == "__main__":
    main()

